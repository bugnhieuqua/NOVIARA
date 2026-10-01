# -*- coding: utf-8 -*-
"""
Script kết xuất toàn bộ 8 sơ đồ CSDL chuyên sâu và nhúng vào:
File/File dữ liệu/Tong_quan_va_Dac_ta_CSDL_NOVIARA.docx
"""
import os
import sys
if sys.stdout:
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass
from pathlib import Path
import urllib.request
import base64
import json
import time

import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

img_dir = Path("File/File dữ liệu/images")
img_dir.mkdir(parents=True, exist_ok=True)
out_dir = Path("File/File dữ liệu")

# Định nghĩa 8 sơ đồ CSDL
CSDL_DIAGRAMS = {
    # 1. Master ERD
    "erd_01_master": {
        "theme": "neutral",
        "code": """erDiagram
    DEPARTMENTS ||--o{ ACCOUNTS : manages
    DEPARTMENTS ||--o{ CLASSES : offers
    ACCOUNTS ||--o{ CLASSES : lectures
    ACCOUNTS ||--o{ GROUPING_SESSIONS : executes
    CLASSES ||--o{ CLASS_STUDENTS : enrolls
    CLASSES ||--o{ GROUPING_SESSIONS : partitions
    CLASSES ||--o{ SURVEY_SUBMISSIONS : collects
    CLASSES ||--o{ UPLOADED_FILES : stores
    STUDENTS ||--o{ CLASS_STUDENTS : belongs_to
    STUDENTS ||--o{ SURVEY_SUBMISSIONS : submits
    STUDENTS ||--o{ GROUP_MEMBERS : participates
    GROUPING_SESSIONS ||--|| SESSION_CONFIGS : configured_by
    GROUPING_SESSIONS ||--o{ GROUPS : produces
    GROUPS ||--o{ GROUP_MEMBERS : contains
    GROUPS ||--o| GROUP_EXPLANATIONS : explained_by

    DEPARTMENTS {
        varchar id PK
        varchar code
        varchar name
    }
    ACCOUNTS {
        varchar id PK
        varchar username
        varchar password_hash
        varchar name
        varchar role
    }
    CLASSES {
        varchar id PK
        varchar code
        varchar name
        varchar department
        varchar lecturer_id FK
        boolean is_survey_active
        varchar survey_title
    }
    STUDENTS {
        varchar student_id PK
        varchar name
        varchar email
        varchar phone
        varchar gender
        real gpa
        varchar primary_skill
        varchar disc_dominant
        boolean is_leader_candidate
    }
    CLASS_STUDENTS {
        int id PK
        varchar class_id FK
        varchar student_id FK
        real gpa
        real skill_frontend
        real skill_backend
        real skill_database
        varchar source
    }
    SURVEY_SUBMISSIONS {
        varchar id PK
        varchar class_id FK
        varchar student_id FK
        real gpa
        varchar primary_skill
        varchar disc_dominant
        boolean is_leader_candidate
        datetime submitted_at
    }
    GROUPING_SESSIONS {
        varchar id PK
        varchar class_id FK
        varchar title
        varchar status
        int total_students
        real overall_fitness
        int execution_time_ms
    }
    SESSION_CONFIGS {
        varchar session_id PK
        int target_group_count
        int min_members
        int max_members
        real weight_skill_balance
        real weight_disc_diversity
        real weight_gpa_balance
        boolean require_leader
    }
    GROUPS {
        varchar id PK
        varchar session_id FK
        int group_number
        varchar name
        varchar leader_id FK
        real avg_gpa
        real compatibility_score
    }
    GROUP_MEMBERS {
        int id PK
        varchar group_id FK
        varchar student_id FK
        boolean is_leader
        varchar assigned_role
    }
    GROUP_EXPLANATIONS {
        int id PK
        varchar group_id FK
        text summary
        text leadership_analysis
        text disc_synergy
        text ai_pedagogical_advice
    }
    UPLOADED_FILES {
        varchar id PK
        varchar filename
        int file_size
        int total_records
        varchar class_id FK
    }
"""
    },

    # 2. Sơ đồ Kiến trúc Lưu trữ SQLite WAL
    "erd_02_wal_architecture": {
        "theme": "neutral",
        "code": """flowchart TD
    subgraph ClientApp [Tang Ung Dung FastAPI]
        FASTAPI[FastAPI Backend Server Python 3.11]
    end

    subgraph PoolLayer [Tang Dieu Phap Ket Noi]
        POOL[Connection Manager<br>PRAGMA foreign_keys = ON]
        FASTAPI --> POOL
    end

    subgraph SQLiteEngine [Co Che Luu Tru SQLite WAL Mode]
        WAL_LOG[(smartgroup.db-wal<br>Write-Ahead Log<br>Doc va Ghi khong bi khoa)]
        SHM_MEM[(smartgroup.db-shm<br>Shared Memory Index)]
        MAIN_DB[(smartgroup.db<br>Tep CSDL 3NF Chinh Thuc)]
    end

    POOL -->|Ghi Transaction moi| WAL_LOG
    POOL -->|Doc Snapshot hien tai| MAIN_DB
    POOL -.->|Tra cuu index nhanh| SHM_MEM
    WAL_LOG -->|Checkpointing dong bo| MAIN_DB
"""
    },

    # 3. Phân hệ 1: Tài khoản & Khoa
    "erd_03_sub_auth": {
        "theme": "neutral",
        "code": """erDiagram
    DEPARTMENTS ||--o{ ACCOUNTS : manages
    DEPARTMENTS ||--o{ CLASSES : offers
    ACCOUNTS ||--o{ CLASSES : lectures

    DEPARTMENTS {
        varchar id PK
        varchar code
        varchar name
        text description
    }
    ACCOUNTS {
        varchar id PK
        varchar username
        varchar password_hash
        varchar name
        varchar email
        varchar role
        bool is_default_password
        bool must_change_password
    }
    CLASSES {
        varchar id PK
        varchar code
        varchar name
        varchar department
        varchar lecturer_id FK
        bool is_survey_active
    }
"""
    },

    # 4. Phân hệ 2: Lớp học & Sinh viên & File tracking
    "erd_04_sub_students": {
        "theme": "neutral",
        "code": """erDiagram
    CLASSES ||--o{ CLASS_STUDENTS : enrolls
    STUDENTS ||--o{ CLASS_STUDENTS : belongs_to
    CLASSES ||--o{ UPLOADED_FILES : tracks

    CLASSES {
        varchar id PK
        varchar code
        varchar name
    }
    STUDENTS {
        varchar student_id PK
        varchar name
        varchar email
        varchar gender
        real gpa
        varchar primary_skill
        varchar disc_dominant
    }
    CLASS_STUDENTS {
        int id PK
        varchar class_id FK
        varchar student_id FK
        real gpa
        real skill_frontend
        real skill_backend
        real skill_database
        varchar source
    }
    UPLOADED_FILES {
        varchar id PK
        varchar filename
        int file_size
        int total_records
        varchar class_id FK
    }
"""
    },

    # 5. Phân hệ 3: Khảo sát DISC trực tuyến
    "erd_05_sub_survey": {
        "theme": "neutral",
        "code": """erDiagram
    CLASSES ||--o{ SURVEY_SUBMISSIONS : collects
    STUDENTS ||--o{ SURVEY_SUBMISSIONS : submits

    CLASSES {
        varchar id PK
        varchar name
        bool is_survey_active
        varchar survey_title
    }
    STUDENTS {
        varchar student_id PK
        varchar name
        varchar disc_dominant
        bool is_leader_candidate
    }
    SURVEY_SUBMISSIONS {
        varchar id PK
        varchar class_id FK
        varchar student_id FK
        real gpa
        varchar primary_skill
        varchar disc_dominant
        bool is_leader_candidate
        text answers_json
        datetime submitted_at
    }
"""
    },

    # 6. Phân hệ 4: GA, Nhóm & Lời khuyên AI
    "erd_06_sub_ga": {
        "theme": "neutral",
        "code": """erDiagram
    CLASSES ||--o{ GROUPING_SESSIONS : partitions
    GROUPING_SESSIONS ||--|| SESSION_CONFIGS : configures
    GROUPING_SESSIONS ||--o{ GROUPS : produces
    GROUPS ||--o{ GROUP_MEMBERS : contains
    STUDENTS ||--o{ GROUP_MEMBERS : assigns
    GROUPS ||--o| GROUP_EXPLANATIONS : explains

    GROUPING_SESSIONS {
        varchar id PK
        varchar class_id FK
        varchar status
        int total_students
        real overall_fitness
    }
    SESSION_CONFIGS {
        varchar session_id PK
        int target_group_count
        real weight_skill_balance
        real weight_disc_diversity
        real weight_gpa_balance
    }
    GROUPS {
        varchar id PK
        varchar session_id FK
        int group_number
        varchar leader_id FK
        real avg_gpa
        real compatibility_score
    }
    GROUP_MEMBERS {
        int id PK
        varchar group_id FK
        varchar student_id FK
        bool is_leader
    }
    GROUP_EXPLANATIONS {
        int id PK
        varchar group_id FK
        text summary
        text leadership_analysis
        text disc_synergy
        text ai_pedagogical_advice
    }
    STUDENTS {
        varchar student_id PK
        varchar name
    }
"""
    },

    # 7. Sơ đồ Mạng lưới Khóa Ngoại
    "erd_07_fk_network": {
        "theme": "neutral",
        "code": """flowchart TD
    ACCOUNTS[(accounts)]
    CLASSES[(classes)]
    STUDENTS[(students)]
    CLASS_STUDENTS[(class_students)]
    SURVEY_SUBMISSIONS[(survey_submissions)]
    GROUPING_SESSIONS[(grouping_sessions)]
    SESSION_CONFIGS[(session_configs)]
    GROUPS[(groups)]
    GROUP_MEMBERS[(group_members)]
    GROUP_EXPLANATIONS[(group_explanations)]
    UPLOADED_FILES[(uploaded_files)]

    ACCOUNTS -->|RESTRICT| CLASSES
    CLASSES -->|ON DELETE CASCADE| CLASS_STUDENTS
    STUDENTS -->|ON UPDATE CASCADE| CLASS_STUDENTS
    CLASSES -->|ON DELETE CASCADE| SURVEY_SUBMISSIONS
    STUDENTS -->|ON UPDATE CASCADE| SURVEY_SUBMISSIONS
    CLASSES -->|SET NULL| GROUPING_SESSIONS
    GROUPING_SESSIONS -->|ON DELETE CASCADE| SESSION_CONFIGS
    GROUPING_SESSIONS -->|ON DELETE CASCADE| GROUPS
    GROUPS -->|ON DELETE CASCADE| GROUP_MEMBERS
    STUDENTS -->|RESTRICT| GROUP_MEMBERS
    GROUPS -->|ON DELETE CASCADE| GROUP_EXPLANATIONS
    CLASSES -->|SET NULL| UPLOADED_FILES
"""
    },

    # 8. Sơ đồ Luồng Dữ Liệu CSDL
    "erd_08_dataflow": {
        "theme": "neutral",
        "code": """flowchart TD
    subgraph ExcelInput [1. Nhap Lieu Excel & File Tracking]
        F1[File Excel danh sach SV] -->|Upload BLOB| UF[(uploaded_files)]
        UF -->|Parse OpenPyXL| STU[(students)]
        UF -->|Ghi danh Lop| CS[(class_students)]
    end

    subgraph SurveyInput [2. Khao Sat Truc Tuyen DISC]
        SV[Sinh vien lam bai] -->|Submit 12 cau DISC| SS[(survey_submissions)]
        SS -->|Dong bo profile| STU
    end

    subgraph GAEngine [3. Thuat Toan Di Truyen GA]
        CS -->|Trich xuat ky nang| GA[Loi Thuat Toan GA]
        SS -->|Trich xuat DISC| GA
        CFG[(session_configs)] -->|Tham so & Trong so| GA
        GA -->|Khoi tao phien nhap| GS[(grouping_sessions)]
        GA -->|Toi uu K nhom| GRP[(groups)]
        GA -->|Bo nhiem Leader & Thanh vien| GM[(group_members)]
    end

    subgraph AIAdvisor [4. Tro Ly Su Pham AI]
        GRP -->|Snapshot nhom| AI[Google Gemini AI]
        GM -->|Doi ngu & Leader| AI
        AI -->|Phan tich & Loi khuyen| GE[(group_explanations)]
    end

    subgraph Output [5. Cong Bo & Tra Cuu Truc Tuyen]
        GS -->|status: published| PUB[Cong Cong Bo Ket Qua]
        PUB -->|Nhap MSSV| LOOKUP[Sinh Vien Tra Cuu Nhom]
    end
"""
    }
}

def render_diagrams():
    print(f"Bat dau ket xuat {len(CSDL_DIAGRAMS)} so do CSDL...")
    for idx, (name, config) in enumerate(CSDL_DIAGRAMS.items(), 1):
        target = img_dir / f"{name}.png"
        print(f"[{idx}/{len(CSDL_DIAGRAMS)}] Dang ket xuat: {name}...", end=" ")
        try:
            state = {
                "code": config["code"].strip(),
                "mermaid": {"theme": config.get("theme", "neutral")}
            }
            b64 = base64.b64encode(json.dumps(state).encode("utf-8")).decode("ascii")
            url = "https://mermaid.ink/img/" + b64
            req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
            with urllib.request.urlopen(req, timeout=25) as resp:
                data = resp.read()
                with open(target, "wb") as f:
                    f.write(data)
            print(f"THANH CONG ({len(data):,} bytes)")
            time.sleep(0.4)
        except Exception as e:
            print(f"LOI: {e}")

# Các hàm phụ trợ định dạng bảng DOCX
def set_cell_background(cell, color_hex):
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="cbd5e1"):
    tblPr = table._tbl.tblPr
    tblBorders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="6" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(tblBorders)

def add_header(doc, title, subtitle):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    run_title = p.add_run(title)
    run_title.font.name = 'Arial'
    run_title.font.size = Pt(21)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(15, 23, 42)

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(16)
    run_sub = p_sub.add_run(subtitle)
    run_sub.font.name = 'Arial'
    run_sub.font.size = Pt(11.5)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(79, 70, 229)

def add_h1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Arial'
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = RGBColor(30, 41, 59)

def add_h2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Arial'
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.color.rgb = RGBColor(51, 65, 85)

def add_p(doc, text, bold_prefix=""):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Arial'
        r_pre.font.size = Pt(10)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(15, 23, 42)
    run = p.add_run(text)
    run.font.name = 'Arial'
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(30, 41, 59)
    return p

def add_diagram_image(doc, filename, caption, width_inches=6.2):
    file_path = img_dir / filename
    if not file_path.exists():
        print(f"Canh bao: Khong tim thay anh {file_path}")
        return
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(8)
    p_img.paragraph_format.space_after = Pt(4)
    run_img = p_img.add_run()
    run_img.add_picture(str(file_path), width=Inches(width_inches))

    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(2)
    p_cap.paragraph_format.space_after = Pt(12)
    r_cap = p_cap.add_run(caption)
    r_cap.font.name = 'Arial'
    r_cap.font.size = Pt(9.5)
    r_cap.font.italic = True
    r_cap.font.bold = True
    r_cap.font.color.rgb = RGBColor(67, 56, 202)

def format_table(table, col_widths, headers, rows):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table, "cbd5e1")

    # Header
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], "1e293b")
        set_cell_margins(hdr_cells[i], top=100, bottom=100, left=120, right=120)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.name = 'Arial'
            r.font.size = Pt(9)
            r.font.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)

    # Body
    for r_idx, row_data in enumerate(rows):
        row_cells = table.rows[r_idx + 1].cells
        bg_color = "f8fafc" if r_idx % 2 == 1 else "ffffff"
        for c_idx, val in enumerate(row_data):
            row_cells[c_idx].text = str(val)
            set_cell_background(row_cells[c_idx], bg_color)
            set_cell_margins(row_cells[c_idx], top=70, bottom=70, left=100, right=100)
            p = row_cells[c_idx].paragraphs[0]
            if len(col_widths) > 2 and c_idx in [0, 2, 3, 4]:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.font.name = 'Arial'
                r.font.size = Pt(8.5)
                r.font.color.rgb = RGBColor(30, 41, 59)

    for row in table.rows:
        for c_idx, w in enumerate(col_widths):
            row.cells[c_idx].width = Inches(w)

def build_full_csdl_document():
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0